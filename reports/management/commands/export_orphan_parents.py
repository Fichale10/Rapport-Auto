"""Liste les Site Parent 1 / 2 qui ne correspondent à aucun site existant,
pour que l'équipe terrain corrige la base source (cause du mauvais
positionnement dans l'onglet Architecture globale)."""

import openpyxl
from django.core.management.base import BaseCommand
from openpyxl.styles import Alignment, Font, PatternFill

from reports.models import Site


class Command(BaseCommand):
    help = "Exporte un Excel des références Site Parent 1/2 orphelines (site introuvable)."

    def add_arguments(self, parser):
        parser.add_argument(
            'output', nargs='?', default='orphan_parents.xlsx',
            type=str, help='Chemin du fichier Excel à générer')

    def handle(self, *args, **options):
        rows = list(Site.objects.values_list(
            'site_name', 'site_parent_1', 'site_parent_2', 'region', 'base'))
        names = {(n or '').strip().upper() for n, _, _, _, _ in rows}

        orphans = []
        for name, p1, p2, region, base in rows:
            for field, parent in (('Site Parent 1', p1), ('Site Parent 2', p2)):
                parent = (parent or '').strip()
                if parent and parent.upper() not in names:
                    orphans.append((name, region, base, field, parent))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Références orphelines'
        headers = ['Nom du site', 'Région', 'Base', 'Champ', 'Valeur introuvable (à corriger)']
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = PatternFill('solid', fgColor='003087')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center')
        ws.freeze_panes = 'A2'

        for name, region, base, field, parent in sorted(orphans, key=lambda r: (r[1] or '', r[0])):
            ws.append([name, region, base, field, parent])

        for col, width in zip('ABCDE', (28, 18, 18, 16, 32)):
            ws.column_dimensions[col].width = width

        wb.save(options['output'])
        self.stdout.write(self.style.SUCCESS(
            f"{len(orphans)} référence(s) orpheline(s) exportée(s) vers {options['output']}"))
