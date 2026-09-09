from datetime import date, datetime, timedelta, timezone
from unittest.mock import patch

import pandas as pd
from pptx import Presentation
from django.test import SimpleTestCase, TestCase

from .dr2_analytics import compute
from .dr2_availability import (
	compute_dr2_sites, filter_availability_day, normalize_dr2_datetimes,
)
from .dr2_meeting_pptx import (
	_causes_breakdown, _dr2_dataset, _efficiency_summary_rows, _gdi_reporting_periods,
	_monthly_comparison, generate_gdi_daily, generate_reunion_hebdo,
)
from .models import Dr2ProcessedDate, Dr2ViolationRecord, Site


class Dr2AvailabilityTests(SimpleTestCase):
	def test_selected_day_excludes_adjacent_hours(self):
		periods = pd.to_datetime([
			'2026-08-19 22:00', '2026-08-19 23:00',
			'2026-08-20 00:00', '2026-08-20 01:00',
		])
		source = pd.DataFrame({
			'period': periods,
			'site_raw': ['SITE_A'] * 4,
			'site_key': ['SITEA'] * 4,
			'value': [0] * 4,
		})

		selected = filter_availability_day(source, date(2026, 8, 20), '2G')

		self.assertEqual(len(selected), 2)
		self.assertEqual(compute_dr2_sites(selected, selected), [])

	def test_normalize_dr2_datetimes_repairs_day_month_inversion(self):
		alarm = datetime(2026, 5, 8, 23, 0, tzinfo=timezone.utc)
		cancel = datetime(2026, 6, 8, 2, 0, tzinfo=timezone.utc)

		normalized_alarm, normalized_cancel = normalize_dr2_datetimes(
			date(2026, 8, 6), alarm, cancel)

		self.assertEqual(
			normalized_alarm,
			datetime(2026, 8, 5, 23, 0, tzinfo=timezone.utc),
		)
		self.assertEqual(
			normalized_cancel,
			datetime(2026, 8, 6, 2, 0, tzinfo=timezone.utc),
		)
		self.assertEqual(normalized_cancel - normalized_alarm, timedelta(hours=3))

	def test_normalize_dr2_datetimes_preserves_legitimate_long_outage(self):
		alarm = datetime(2026, 7, 1, 8, 0, tzinfo=timezone.utc)
		cancel = datetime(2026, 7, 11, 8, 0, tzinfo=timezone.utc)

		self.assertEqual(
			normalize_dr2_datetimes(date(2026, 7, 11), alarm, cancel),
			(alarm, cancel),
		)

	def test_normalize_dr2_datetimes_rejects_distant_swap(self):
		alarm = datetime(2026, 7, 8, 8, 0, tzinfo=timezone.utc)

		self.assertEqual(
			normalize_dr2_datetimes(date(2026, 8, 31), alarm, None),
			(alarm, None),
		)

	def test_normalize_dr2_datetimes_does_not_move_alarm_after_report_day(self):
		alarm = datetime(2026, 7, 8, 8, 0, tzinfo=timezone.utc)

		self.assertEqual(
			normalize_dr2_datetimes(date(2026, 8, 6), alarm, None),
			(alarm, None),
		)

	def test_gdi_reporting_periods_use_last_completed_weekend(self):
		periods = _gdi_reporting_periods(date(2026, 8, 26))

		self.assertEqual(periods['month_start'], date(2026, 8, 1))
		self.assertEqual(periods['detail_start'], date(2026, 8, 21))
		self.assertEqual(periods['detail_end'], date(2026, 8, 23))
		self.assertEqual(periods['meeting_day'], date(2026, 8, 24))


class Dr2AnalyticsTests(TestCase):
	def test_efficiency_summary_uses_incident_duration_and_dr2_counts(self):
		mobile_df = pd.DataFrame([
			{'date': '2026-09-02', 'region': 'Lomé', 'base': 'NOTSE', 'duration_sec': 1800},
			{'date': '2026-09-03', 'region': 'LOME', 'base': 'NOTSE', 'duration_sec': 3600},
			{'date': '2026-09-04', 'region': 'KARA', 'base': 'MANGO', 'duration_sec': 7200},
			{'date': '2026-08-31', 'region': 'LOME', 'base': 'NOTSE', 'duration_sec': 99999},
		])
		region_rows = [
			{'region': 'LOME', 'dr2': 1},
			{'region': 'KARA', 'dr2': 0},
		]
		base_rows = [
			{'base': 'NOTSE', 'dr2': 1},
			{'base': 'MANGO', 'dr2': 0},
		]

		regions, bases = _efficiency_summary_rows(
			mobile_df, region_rows, base_rows,
			date(2026, 9, 1), date(2026, 9, 6),
		)

		region_by_name = {row['label']: row for row in regions}
		base_by_name = {row['label']: row for row in bases}
		self.assertEqual(
			region_by_name['LOME'],
			{'label': 'LOME', 'incidents': 2, 'mttr': '0:45:00', 'dr2': 1, 'efficiency': 50},
		)
		self.assertEqual(
			base_by_name['NOTSE'],
			{'label': 'NOTSE', 'incidents': 2, 'mttr': '0:45:00', 'dr2': 1, 'efficiency': 50},
		)
		self.assertEqual(base_by_name['MANGO']['mttr'], '2:00:00')
		self.assertEqual(base_by_name['MANGO']['efficiency'], 100)

	def test_gdi_deck_places_red_definitions_before_summary(self):
		with patch(
			'reports.dr2_meeting_pptx._mobile_incident_dataframe',
			return_value=pd.DataFrame(columns=[
				'date', 'region', 'base', 'site', 'duration_sec', 'cause',
			]),
		):
			presentation = Presentation(generate_gdi_daily(
				date(2026, 9, 1), date(2026, 9, 6), '07/09/2026',
			))

		definition_slide = presentation.slides[2]
		summary_slide = presentation.slides[3]
		definition_text = ' '.join(
			shape.text for shape in definition_slide.shapes
			if hasattr(shape, 'text')
		)
		summary_text = ' '.join(
			shape.text for shape in summary_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('Définition DR1/DR2', definition_text)
		self.assertIn('SYNTHÈSE EXÉCUTIVE DR2', summary_text)

		table = next(shape.table for shape in definition_slide.shapes if shape.has_table)
		self.assertEqual(str(table.cell(0, 0).fill.fore_color.rgb), 'F8696B')
		self.assertEqual(str(table.cell(1, 0).fill.fore_color.rgb), 'FFD1D1')
		self.assertEqual(str(table.cell(2, 0).fill.fore_color.rgb), 'FFE6E6')
		self.assertEqual(
			str(table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.color.rgb),
			'003087',
		)

	def test_gdi_trend_shows_month_to_date_and_frames_selected_period(self):
		for day, count in (
			(date(2026, 9, 4), 2),
			(date(2026, 9, 5), 7),
			(date(2026, 9, 6), 6),
		):
			Dr2ProcessedDate.objects.create(date=day, sites_count=count)
			for index in range(count):
				site_name = (
					'RECURRING-SITE' if index == 0
					else f'SITE-{day.day}-{index}'
				)
				point_bloquant = ''
				if day == date(2026, 9, 4) and index == 0:
					point_bloquant = 'MANQUE DE PDR'
				elif day == date(2026, 9, 4) and index == 1:
					point_bloquant = "PROBLEME D ' ACCES"
				elif day == date(2026, 9, 5) and index < 5:
					point_bloquant = 'Manque de PDR'
				elif day == date(2026, 9, 6) and index < 4:
					point_bloquant = 'MANQUE DE PDR'
				Dr2ViolationRecord.objects.create(
					date=day,
					site_name=site_name,
					region='LOME', categorie='ENERGIE', cause='PANNE',
					point_bloquant=point_bloquant,
					hours_down=4, is_resolved=True,
				)
				Site.objects.get_or_create(
					site_name=site_name,
					defaults={
						'region': 'LOME',
						'base': 'ELAVAGNON' if index == 0 else 'MANGO',
					},
				)
		Site.objects.create(
			site_name='SITE-WITHOUT-DR2', region='PLATEAUX', base='KPALIME',
		)

		mobile_df = pd.DataFrame([
			{
				'date': '2026-09-05', 'region': 'LOME', 'base': 'MANGO',
				'site': f'MOBILE-{index}', 'duration_sec': 3600, 'cause': 'PANNE',
			}
			for index in range(20)
		])
		with patch(
			'reports.dr2_meeting_pptx._mobile_incident_dataframe',
			return_value=mobile_df,
		):
			presentation = Presentation(generate_gdi_daily(
				date(2026, 9, 4), date(2026, 9, 6), '07/09/2026',
			))

		trend_slide = presentation.slides[7]
		chart = next(shape.chart for shape in trend_slide.shapes if shape.has_chart)
		self.assertEqual(
			[category.label for category in chart.plots[0].categories],
			['1', '2', '3', '4', '5', '6'],
		)
		self.assertEqual(chart.series[0].values, (0.0, 0.0, 0.0, 2.0, 7.0, 6.0))
		trend_text = ' '.join(
			shape.text for shape in trend_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('TREND DR2 SEPTEMBRE 2026', trend_text)
		self.assertIn('TDR2 = 15', trend_text)
		self.assertIn('MOY = 5,0', trend_text)
		self.assertGreaterEqual(trend_slide.element.xml.count('val="dash"'), 2)
		self.assertIn('00B050', trend_slide.element.xml)

		detail_slides = [presentation.slides[index] for index in range(8, 11)]
		self.assertEqual(len(detail_slides), 3)
		for detail_slide in detail_slides:
			detail_text = ' '.join(
				shape.text for shape in detail_slide.shapes
				if hasattr(shape, 'text')
			)
			self.assertIn('DETAILS DR2', detail_text)
			self.assertNotIn('VENDREDI', detail_text)
			self.assertNotIn('SAMEDI', detail_text)
			self.assertNotIn('DIMANCHE', detail_text)
			table = next(shape.table for shape in detail_slide.shapes if shape.has_table)
			self.assertEqual(str(table.cell(1, 0).fill.fore_color.rgb), '70AD47')
			self.assertEqual(str(table.cell(1, 2).fill.fore_color.rgb), 'FFC72C')
			self.assertEqual(str(table.cell(1, 13).fill.fore_color.rgb), 'FF0000')

		top_sites_slide = presentation.slides[12]
		top_sites_text = ' '.join(
			shape.text for shape in top_sites_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('TOP SITE OCCURRENCE DR2', top_sites_text)
		self.assertNotIn('CLASSEMENT DES SITES RÉCURRENTS DR2', top_sites_text)
		top_chart = next(shape.chart for shape in top_sites_slide.shapes if shape.has_chart)
		self.assertEqual(top_chart.value_axis.major_unit, 1.0)
		self.assertEqual(top_chart.value_axis.maximum_scale, 4.0)
		self.assertEqual(top_chart.value_axis.tick_labels.number_format, '0')
		categories = [category.label for category in top_chart.plots[0].categories]
		self.assertIn('RECURRING-SITE', categories)
		self.assertFalse(any(category[:2].isdigit() for category in categories))
		point_colors = {
			str(point.format.fill.fore_color.rgb)
			for point in top_chart.series[0].points
		}
		self.assertEqual(point_colors, {'FF0000', 'FFC72C'})

		region_base_slide = presentation.slides[13]
		region_base_text = ' '.join(
			shape.text for shape in region_base_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('DR2 COUNT BY BASE TECHNIQUE', region_base_text)
		region_table = next(
			shape.table for shape in region_base_slide.shapes if shape.has_table
		)
		self.assertEqual(
			region_table.cell(0, 0).text, 'VIOLATION DR2 / REGIONS',
		)
		self.assertEqual(
			[region_table.cell(2, column).text for column in range(4)],
			['LOME', '372', '30', '15'],
		)
		self.assertEqual(region_table.cell(8, 2).text, '90')
		base_chart = next(
			shape.chart for shape in region_base_slide.shapes if shape.has_chart
		)
		base_categories = [
			category.label for category in base_chart.plots[0].categories
		]
		series_by_name = {series.name: series for series in base_chart.series}
		actual_by_base = dict(zip(
			base_categories, series_by_name['DR2 COUNT'].values,
		))
		target_by_base = dict(zip(
			base_categories, series_by_name['TARGET DR2'].values,
		))
		self.assertEqual(
			actual_by_base,
			{'KPALIME': 0.0, 'MANGO': 12.0, 'ELAVAGNON': 3.0},
		)
		self.assertEqual(
			target_by_base,
			{'KPALIME': 1.0, 'MANGO': 1.0, 'ELAVAGNON': 1.0},
		)

		efficiency_slide = presentation.slides[14]
		efficiency_text = ' '.join(
			shape.text for shape in efficiency_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('EFFICACITE DR2', efficiency_text)
		efficiency_tables = [
			shape.table for shape in efficiency_slide.shapes if shape.has_table
		]
		self.assertEqual(len(efficiency_tables), 2)
		self.assertEqual(
			[efficiency_tables[0].cell(0, column).text for column in range(5)],
			['REGION', 'Nbr I', 'MTTR', 'DR2', 'EFFICACITE DR2'],
		)
		self.assertEqual(
			[efficiency_tables[1].cell(0, column).text for column in range(5)],
			['BASE TECH', "NBRE D'INCIDENT", 'MTTR INC', 'NBRE DE DR2', 'Efficacité DR2'],
		)
		region_values = {
			row.cells[0].text: [cell.text for cell in row.cells]
			for row in list(efficiency_tables[0].rows)[1:]
		}
		base_values = {
			row.cells[0].text: [cell.text for cell in row.cells]
			for row in list(efficiency_tables[1].rows)[1:]
		}
		self.assertEqual(region_values['LOME'], ['LOME', '20', '1:00:00', '15', '25%'])
		self.assertEqual(base_values['MANGO'], ['MANGO', '20', '1:00:00', '12', '40%'])

		blocking_slide = presentation.slides[15]
		blocking_text = ' '.join(
			shape.text for shape in blocking_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('TREND BLOCK POINT', blocking_text)
		self.assertIn('DR2- PB= 4', blocking_text)
		blocking_chart = next(
			shape.chart for shape in blocking_slide.shapes if shape.has_chart
		)
		self.assertEqual(
			[series.name for series in blocking_chart.series],
			[
				'MANQUE DE PDR', "PROBLEME D'ACCES", 'RESPECT NORME HSE',
				'VANDALISME', 'RESPECT NORME SURETE',
				'RESPECT NORME SURETE (ZONE ROUGE)',
			],
		)
		self.assertEqual(len(blocking_chart.series[0].values), 30)
		self.assertEqual(
			blocking_chart.series[0].values[3:6], (1.0, 5.0, 4.0),
		)
		blocking_table = next(
			shape.table for shape in blocking_slide.shapes if shape.has_table
		)
		self.assertEqual(blocking_table.cell(0, 0).text, 'POINTS BLOQUANTS')
		self.assertEqual(
			[blocking_table.cell(row, 1).text for row in range(2, 8)],
			['10', '1', '0', '0', '0', '0'],
		)
		self.assertEqual(str(blocking_table.cell(2, 1).fill.fore_color.rgb), 'F8696B')
		self.assertEqual(str(blocking_table.cell(3, 1).fill.fore_color.rgb), 'FFE680')
		self.assertEqual(str(blocking_table.cell(4, 1).fill.fore_color.rgb), '63BE7B')
		self.assertEqual(blocking_table.cell(8, 1).text, '11')

		matrix_slide = presentation.slides[17]
		matrix_text = ' '.join(
			shape.text for shape in matrix_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('TABLEAU DE CROISEMENT REGION / METIERS', matrix_text)
		matrix = next(shape.table for shape in matrix_slide.shapes if shape.has_table)
		self.assertEqual(len(matrix.rows), 9)
		self.assertEqual(len(matrix.columns), 14)
		self.assertEqual(matrix.cell(0, 0).text, 'TABLEAU DE CROISEMENT')
		self.assertTrue(matrix.cell(0, 0).is_merge_origin)
		self.assertEqual(str(matrix.cell(0, 0).fill.fore_color.rgb), 'FFC72C')
		self.assertEqual(matrix.cell(1, 2).text, '')
		self.assertTrue(matrix.cell(1, 2).is_merge_origin)
		escalation_label = next(
			shape for shape in matrix_slide.shapes
			if getattr(shape, 'text', '') == 'BY ESCALADE'
		)
		self.assertEqual(escalation_label.rotation, 270.0)
		self.assertEqual(
			[matrix.cell(row, 0).text for row in range(2, 8)],
			['LOME', 'MARITIME', 'PLATEAUX', 'CENTRALE', 'KARA', 'SAVANES'],
		)
		self.assertEqual(
			[matrix.cell(1, column).text for column in range(3, 14)],
			[
				'ENERGIE', 'RAN-FIELD O', 'TRANS FH-FIELD O', 'TRANS IP',
				'TRANS FO', 'TRANS FTTM', 'PROJET', 'BSS', 'ENVIRONNEMENT',
				'INFRA', 'ENERGIE / TRANS / RAN',
			],
		)
		self.assertEqual(str(matrix.cell(2, 3).fill.fore_color.rgb), 'F8696B')
		self.assertEqual(str(matrix.cell(3, 3).fill.fore_color.rgb), '63BE7B')
		self.assertEqual(str(matrix.cell(8, 0).fill.fore_color.rgb), 'F4B183')
		self.assertEqual(str(matrix.cell(8, 3).fill.fore_color.rgb), 'D9D9D9')

	def test_weekly_deck_keeps_dense_tail_and_summary_separate(self):
		for day, count in (
			(date(2026, 8, 21), 4),
			(date(2026, 8, 22), 4),
			(date(2026, 8, 23), 8),
		):
			Dr2ProcessedDate.objects.create(date=day, sites_count=count)
			for index in range(count):
				Dr2ViolationRecord.objects.create(
					date=day, site_name=f'SITE-{day.day}-{index}',
					region='LOME', categorie='ENERGIE', cause='PANNE',
					hours_down=4, is_resolved=True,
				)

		with patch('reports.dr2_meeting_pptx._dr1_violations', return_value=[]):
			presentation = Presentation(generate_reunion_hebdo(
				date(2026, 8, 21), date(2026, 8, 23), '28/08/2026',
			))

		self.assertEqual(len(presentation.slides), 10)
		self.assertEqual(
			sum(shape.has_table for shape in presentation.slides[6].shapes), 1,
		)
		detail_text = ' '.join(
			shape.text for shape in presentation.slides[6].shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('DETAILS DR2', detail_text)
		self.assertNotIn('VENDREDI', detail_text)
		self.assertNotIn('SAMEDI', detail_text)
		self.assertNotIn('DIMANCHE', detail_text)
		self.assertEqual(
			sum(shape.has_table for shape in presentation.slides[7].shapes), 2,
		)

	def test_weekly_deck_uses_month_trend_and_groups_sparse_tail(self):
		plan = (
			(date(2026, 8, 24), 6),
			(date(2026, 8, 25), 3),
			(date(2026, 8, 26), 2),
			(date(2026, 8, 27), 1),
		)
		Dr2ProcessedDate.objects.create(date=date(2026, 8, 4), sites_count=1)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 4), site_name='MONTH-SITE', region='LOME',
			categorie='ENERGIE', cause='PANNE', hours_down=4,
		)
		for day, count in plan:
			Dr2ProcessedDate.objects.create(date=day, sites_count=count)
			for index in range(count):
				Dr2ViolationRecord.objects.create(
					date=day, site_name=f'SITE-{day.day}-{index}',
					region='LOME', categorie='ENERGIE', cause='PANNE',
					hours_down=4, is_resolved=True,
				)

		with patch('reports.dr2_meeting_pptx._dr1_violations', return_value=[]):
			presentation = Presentation(generate_reunion_hebdo(
				date(2026, 8, 24), date(2026, 8, 27), '28/08/2026',
			))

		self.assertEqual(len(presentation.slides), 9)
		trend_text = ' '.join(
			shape.text for shape in presentation.slides[3].shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('13 DR2', trend_text)
		self.assertIn('Nbr DR2 = 12', trend_text)
		combined_slide = presentation.slides[6]
		self.assertEqual(sum(shape.has_table for shape in combined_slide.shapes), 4)
		combined_text = ' '.join(
			cell.text
			for shape in combined_slide.shapes if shape.has_table
			for row in shape.table.rows for cell in row.cells
		)
		combined_labels = ' '.join(
			shape.text for shape in combined_slide.shapes
			if hasattr(shape, 'text')
		)
		self.assertIn('26-08-2026', combined_labels)
		self.assertIn('27-08-2026', combined_labels)
		self.assertIn('TOTAL DR2 = 12', combined_text)

	def test_gdi_cause_percentages_include_records_without_cause(self):
		day = date(2026, 8, 23)
		for site, cause in (
			('SITE-A', 'RADIO ODU HS'),
			('SITE-B', 'RADIO ODU HS'),
			('SITE-C', ''),
			('SITE-D', ''),
		):
			Dr2ViolationRecord.objects.create(
				date=day, site_name=site, cause=cause, hours_down=3,
			)

		result = _causes_breakdown(
			Dr2ViolationRecord.objects.filter(date=day), total=4,
		)

		self.assertEqual(result, [('RADIO ODU HS', 2, 50)])

	def test_gdi_dataset_uses_processed_days_and_selected_period_end(self):
		for day in (date(2026, 8, 21), date(2026, 8, 22), date(2026, 8, 23)):
			Dr2ProcessedDate.objects.create(date=day, sites_count=0)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 21), site_name='SITE-A', region='LOME',
			categorie='ENERGIE', hours_down=3,
		)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 23), site_name='SITE-B', region='KARA',
			categorie='RAN-FIELD O', hours_down=4,
		)

		result = _dr2_dataset(date(2026, 8, 1), date(2026, 8, 23))

		self.assertEqual(result['period_days'], 3)
		self.assertEqual(result['moyenne'], 0.67)
		self.assertEqual(result['nbre_j1'], 1)

	def test_monthly_comparison_keeps_unprocessed_months_unavailable(self):
		Dr2ProcessedDate.objects.create(date=date(2026, 8, 1), sites_count=1)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 1), site_name='SITE-A', region='LOME',
			categorie='ENERGIE', hours_down=3,
		)

		months = _monthly_comparison(date(2026, 8, 23))

		self.assertEqual([month['moyenne'] for month in months], [None, None, 1.0])
		self.assertEqual([month['processed_days'] for month in months], [0, 0, 1])

	def test_average_and_trend_use_processed_days(self):
		for day, count in ((date(2026, 8, 10), 2), (date(2026, 8, 11), 1),
						   (date(2026, 8, 12), 0)):
			Dr2ProcessedDate.objects.create(date=day, sites_count=count)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 10), site_name='SITE-A', region='LOME',
			categorie='ENERGIE', cause='PANNE', hours_down=3, is_resolved=True,
		)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 10), site_name='SITE-B', region='KARA',
			categorie='RAN', cause='LIEN', hours_down=4, is_resolved=False,
		)
		Dr2ViolationRecord.objects.create(
			date=date(2026, 8, 11), site_name='SITE-A', region='LOME',
			categorie='ENERGIE', cause='PANNE', hours_down=5, is_resolved=False,
		)

		result = compute(date(2026, 8, 10), date(2026, 8, 13))

		self.assertEqual(result['kpi']['violations'], 3)
		self.assertEqual(result['kpi']['processed_days'], 3)
		self.assertEqual(result['kpi']['average'], 1.0)
		self.assertEqual(result['kpi']['recurrent_sites'], 1)
		self.assertEqual([point['count'] for point in result['trend']], [2, 1, 0, None])
