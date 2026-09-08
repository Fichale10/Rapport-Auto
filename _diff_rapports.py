import openpyxl
import sys

f1 = r"RAPPORT JOURNALIER 02-08-2026 (1).xlsx"  # dev
f2 = r"RAPPORT JOURNALIER 02-08-2026 (2).xlsx"  # prod

wb1 = openpyxl.load_workbook(f1, data_only=True)
wb2 = openpyxl.load_workbook(f2, data_only=True)

print("Sheets DEV :", wb1.sheetnames)
print("Sheets PROD:", wb2.sheetnames)
print()

common = [s for s in wb1.sheetnames if s in wb2.sheetnames]
only1 = [s for s in wb1.sheetnames if s not in wb2.sheetnames]
only2 = [s for s in wb2.sheetnames if s not in wb1.sheetnames]
if only1:
    print("Feuilles seulement DEV :", only1)
if only2:
    print("Feuilles seulement PROD:", only2)
print()

for sheet in common:
    ws1 = wb1[sheet]
    ws2 = wb2[sheet]
    max_row = max(ws1.max_row, ws2.max_row)
    max_col = max(ws1.max_column, ws2.max_column)
    diffs = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v1 = ws1.cell(row=r, column=c).value
            v2 = ws2.cell(row=r, column=c).value
            if v1 != v2:
                diffs.append((r, c, v1, v2))
    print(f"=== Feuille '{sheet}' : {len(diffs)} differences (dims dev={ws1.max_row}x{ws1.max_column}, prod={ws2.max_row}x{ws2.max_column}) ===")
    for r, c, v1, v2 in diffs[:60]:
        col_letter = openpyxl.utils.get_column_letter(c)
        print(f"  {col_letter}{r}: DEV={v1!r}  PROD={v2!r}")
    if len(diffs) > 60:
        print(f"  ... et {len(diffs)-60} autres differences")
    print()
